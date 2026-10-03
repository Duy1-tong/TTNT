"""Module xuat bao cao ra file Excel (.xlsx) bang openpyxl."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def xuat_bang_du_lieu_ra_excel(
    duong_dan_file: Path,
    tieu_de_bao_cao: str,
    ten_cac_cot: Sequence[str],
    du_lieu_cac_hang: Sequence[Sequence[Any]],
) -> Path:
    """Xuat mot bang du lieu (tieu de + danh sach hang) ra file Excel co dinh dang.

    Tra ve duong dan file da tao. Ham nay tu tao thu muc chua neu chua co.
    """
    duong_dan_file.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    trang_tinh = workbook.active
    trang_tinh.title = "BaoCao"

    # ----- Tieu de bao cao -----
    so_cot = len(ten_cac_cot)
    trang_tinh.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(so_cot, 1))
    o_tieu_de = trang_tinh.cell(row=1, column=1, value=tieu_de_bao_cao)
    o_tieu_de.font = Font(size=14, bold=True, color="FFFFFF")
    o_tieu_de.alignment = Alignment(horizontal="center", vertical="center")
    o_tieu_de.fill = PatternFill(start_color="2E5395", end_color="2E5395", fill_type="solid")
    trang_tinh.row_dimensions[1].height = 28

    # ----- Dong tieu de cot -----
    dong_tieu_de_cot = 3
    for chi_so_cot, ten_cot in enumerate(ten_cac_cot, start=1):
        o = trang_tinh.cell(row=dong_tieu_de_cot, column=chi_so_cot, value=ten_cot)
        o.font = Font(bold=True, color="FFFFFF")
        o.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        o.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    # ----- Du lieu -----
    for chi_so_hang, hang_du_lieu in enumerate(du_lieu_cac_hang, start=dong_tieu_de_cot + 1):
        for chi_so_cot, gia_tri in enumerate(hang_du_lieu, start=1):
            o = trang_tinh.cell(row=chi_so_hang, column=chi_so_cot, value=gia_tri)
            o.alignment = Alignment(horizontal="center", vertical="center")

    # ----- Tu dong dieu chinh do rong cot -----
    for chi_so_cot in range(1, so_cot + 1):
        do_rong_toi_da = len(str(ten_cac_cot[chi_so_cot - 1]))
        for hang_du_lieu in du_lieu_cac_hang:
            if chi_so_cot - 1 < len(hang_du_lieu):
                do_rong_toi_da = max(do_rong_toi_da, len(str(hang_du_lieu[chi_so_cot - 1])))
        trang_tinh.column_dimensions[get_column_letter(chi_so_cot)].width = min(
            max(do_rong_toi_da + 4, 12), 40
        )

    workbook.save(duong_dan_file)
    return duong_dan_file
